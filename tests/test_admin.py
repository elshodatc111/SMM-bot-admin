import io
import re

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from conftest import make_user, make_plan_active


@pytest.fixture()
def client(env, sent):
    from app.main import app
    with TestClient(app, follow_redirects=False) as c:
        yield c


def login(c):
    r = c.post("/admin/login", data={"password": "admin-password-1"})
    assert r.status_code == 303
    return c


def csrf(c, url="/admin/"):
    html = c.get(url).text
    return re.search(r'name="csrf" value="([0-9a-f]+)"', html).group(1)


def test_login_required_and_wrong_password(client):
    assert client.get("/admin/").status_code == 303
    assert client.get("/admin/users").status_code == 303
    assert client.post("/admin/login", data={"password": "x"}).status_code == 401
    assert client.get("/admin/export/users.xlsx").status_code == 303


def test_lockout_after_failures(client):
    for _ in range(5):
        client.post("/admin/login", data={"password": "bad"})
    assert client.post("/admin/login", data={"password": "admin-password-1"}).status_code == 429
    from app.admin import auth
    auth._fails.clear()


def test_csrf_enforced(client):
    login(client)
    assert client.post("/admin/plans/1", data={"title": "x", "days": 30, "price": 1000}).status_code == 403
    assert client.post("/admin/logout", data={"csrf": "wrong"}).status_code == 403


def test_all_pages_render(client):
    from app import licenses
    uid = make_user()
    p = make_plan_active("y1", 960000)
    lid = licenses.issue_manual(uid, p["id"])
    login(client)
    for url in ["/admin/", "/admin/stats", "/admin/stats?days=90", "/admin/users", "/admin/users?qs=ali", f"/admin/users/{uid}",
                "/admin/plans", "/admin/payments", "/admin/payments?status=paid", "/admin/licenses", "/admin/licenses?status=new&qs=TGP",
                f"/admin/licenses/{lid}", "/admin/devices", "/admin/broadcast", "/admin/settings"]:
        r = client.get(url)
        assert r.status_code == 200, url
        assert "Traceback" not in r.text


def test_plan_price_edit(client):
    from app import db
    login(client)
    t = csrf(client, "/admin/plans")
    r = client.post("/admin/plans/1", data={"csrf": t, "title": "1 oylik", "days": 30, "price": "150 000", "active": "1"})
    assert r.status_code == 303
    p = db.one("SELECT * FROM plans WHERE id=1")
    assert p["price_uzs"] == 150000 and p["active"] == 1
    # narx 0 bo'lsa tarif o'chiq qoladi
    client.post("/admin/plans/1", data={"csrf": t, "title": "1 oylik", "days": 30, "price": "0", "active": "1"})
    assert db.one("SELECT active FROM plans WHERE id=1")["active"] == 0


def test_revoke_extend_reset_and_device_approval(client, sent):
    from app import licenses, db
    uid = make_user()
    p = make_plan_active("y1", 960000)
    lid = licenses.issue_manual(uid, p["id"])
    tok = licenses.get_license(lid)["token"]
    licenses.activate(tok, {"machine_guid": "G"})
    ok, _, rid = licenses.request_device_change(uid, lid)
    login(client)
    t = csrf(client, "/admin/devices")
    assert client.post(f"/admin/devices/{rid}/approve", data={"csrf": t}).status_code == 303
    assert db.one("SELECT fp_json FROM licenses WHERE id=?", (lid,))["fp_json"] == ""
    assert any("ruxsat" in m.lower() for _, m in sent)
    client.post(f"/admin/licenses/{lid}/extend", data={"csrf": t, "days": 10})
    client.post(f"/admin/licenses/{lid}/revoke", data={"csrf": t, "note": "qaytarildi"})
    assert db.one("SELECT status FROM licenses WHERE id=?", (lid,))["status"] == "revoked"
    assert licenses.activate(tok, {"machine_guid": "G"})["code"] == "revoked"
    client.post(f"/admin/licenses/{lid}/unrevoke", data={"csrf": t})
    assert licenses.activate(tok, {"machine_guid": "G"})["ok"]


def test_grant_and_block(client, sent):
    from app import db
    uid = make_user()
    pl = make_plan_active("m1", 1000)
    login(client)
    t = csrf(client, f"/admin/users/{uid}")
    r = client.post(f"/admin/users/{uid}/grant", data={"csrf": t, "plan_id": pl["id"], "note": "sovg'a"})
    assert r.status_code == 303 and db.one("SELECT COUNT(*) n FROM licenses")["n"] == 1
    assert any("TGP-" in m for _, m in sent)
    client.post(f"/admin/users/{uid}/block", data={"csrf": t})
    assert db.one("SELECT blocked FROM users")["blocked"] == 1


def test_exports(client):
    from app import licenses, db
    uid = make_user()
    pl = make_plan_active("m1", 1000)
    licenses.issue_manual(uid, pl["id"])
    login(client)
    for kind in ("users", "payments", "licenses", "stats"):
        r = client.get(f"/admin/export/{kind}.xlsx")
        assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.active.max_row >= (1 if kind == "payments" else 2)
    ws = load_workbook(io.BytesIO(client.get("/admin/export/users.xlsx").content)).active
    assert ws["D2"].value == "Ali Valiyev"


def test_broadcast_queue_and_stats(client, monkeypatch):
    import asyncio
    from app import jobs, notify, db
    make_user(1001)
    make_user(1002, "Vali", "+998900000000")
    db.ex("UPDATE users SET bot_blocked=1 WHERE tg_id=1002")
    got = []

    class FakeBot:
        async def send_message(self, tg_id, text):
            got.append((tg_id, text))

    notify.set_bot(FakeBot())
    try:
        bid = jobs.create_broadcast("Salom", "all")
        assert db.one("SELECT total FROM broadcasts WHERE id=?", (bid,))["total"] == 1      # botni bloklagan chetlab o'tiladi
        asyncio.run(jobs.broadcast_tick())
        assert got == [(1001, "Salom")]
        assert db.one("SELECT status, sent FROM broadcasts WHERE id=?", (bid,))["status"] == "done"
    finally:
        notify.set_bot(None)
    from app.admin import stats
    d = stats.daily(7)
    assert len(d) == 7 and d[0]["users"] >= 0
    assert sum(r["users"] for r in d) == 2


def test_plain_password_in_env_gives_clear_message(client, monkeypatch):
    monkeypatch.setenv("ADMIN_PASSWORD_HASH", "oddiyparol")
    r = client.post("/admin/login", data={"password": "oddiyparol"})
    assert "xesh emas" in r.text


def test_lifetime_license_admin_only(client, sent):
    from app import db, licenses
    uid = make_user()
    login(client)
    t = csrf(client, f"/admin/users/{uid}")
    r = client.post(f"/admin/users/{uid}/grant", data={"csrf": t, "plan_id": "lifetime", "note": "do'st"})
    assert r.status_code == 303
    lic = db.one("SELECT * FROM licenses")
    assert lic["plan_title"] == "Umrbod" and licenses.is_lifetime(lic) and lic["source"] == "admin"
    assert any("umrbod" in m.lower() for _, m in sent)
    res = licenses.activate(lic["token"], {"machine_guid": "G"})
    assert res["ok"]
    import json
    p = json.loads(res["payload"])
    assert p["lifetime"] is True and p["expires_at"] > 4000000000
    # muddat o'zgartirilmaydi, tugamaydi
    licenses.extend(lic["id"], 30)
    assert db.one("SELECT expires_at e FROM licenses")["e"] == licenses.LIFETIME_EXP
    assert licenses.eff_status(licenses.get_license(lic["id"])) == "active"
    # sahifalar ochiladi
    assert "Umrbod" in client.get(f"/admin/licenses/{lic['id']}").text
    assert client.get("/admin/licenses").status_code == 200
    # botdagi tariflar ro'yxatida yo'q
    assert all(p["code"] != "lifetime" for p in licenses.plans(active_only=True))
