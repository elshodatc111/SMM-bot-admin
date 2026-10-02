"""Excel (.xlsx) eksport."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .. import licenses
from ..db import q
from ..util import fmt_dt
from . import stats


def _book(title: str, headers: list[str], rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F766E")
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:300]]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 50)
    ws.freeze_panes = "A2"
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def users_xlsx() -> bytes:
    rows = q("SELECT u.*, (SELECT COUNT(*) FROM licenses l WHERE l.user_id=u.id) lic_n, "
             "(SELECT COALESCE(SUM(amount_tiyin),0)/100 FROM orders o WHERE o.user_id=u.id AND o.status='paid') paid_sum "
             "FROM users u ORDER BY u.id")
    return _book("Foydalanuvchilar",
                 ["ID", "Telegram ID", "Username", "Ism-familiya", "Telefon", "Viloyat", "Ro'yxatdan o'tgan", "Holat",
                  "Litsenziyalar", "Jami to'lov (so'm)"],
                 [[r["id"], r["tg_id"], r["username"], r["full_name"], r["phone"], r["region"],
                   fmt_dt(r["registered_at"]) if r["registered"] else "ro'yxatdan o'tmagan",
                   "bloklangan" if r["blocked"] else ("botni bloklagan" if r["bot_blocked"] else "faol"),
                   r["lic_n"], r["paid_sum"]] for r in rows])


def payments_xlsx() -> bytes:
    rows = q("SELECT o.*, u.full_name, u.phone, u.tg_id, "
             "(SELECT payme_id FROM payments p WHERE p.order_id=o.id ORDER BY p.id DESC LIMIT 1) pid, "
             "(SELECT state FROM payments p WHERE p.order_id=o.id ORDER BY p.id DESC LIMIT 1) pstate "
             "FROM orders o JOIN users u ON u.id=o.user_id ORDER BY o.id DESC")
    lab = {"new": "kutilmoqda", "paid": "to'langan", "cancelled": "bekor"}
    return _book("To'lovlar",
                 ["Buyurtma", "Sana", "Foydalanuvchi", "Telefon", "Tarif", "Summa (so'm)", "Holat", "To'langan vaqt",
                  "Payme tranzaksiya", "Payme holati"],
                 [[r["id"], fmt_dt(r["created_at"]), r["full_name"], r["phone"], r["plan_title"],
                   r["amount_tiyin"] // 100, lab[r["status"]], fmt_dt(r["paid_at"]), r["pid"] or "", r["pstate"] if r["pstate"] is not None else ""]
                  for r in rows])


def licenses_xlsx() -> bytes:
    rows = q("SELECT l.*, u.full_name, u.phone FROM licenses l JOIN users u ON u.id=l.user_id ORDER BY l.id DESC")
    lab = {"new": "aktivlashtirilmagan", "active": "faol", "expired": "muddati tugagan", "revoked": "bekor qilingan"}
    return _book("Litsenziyalar",
                 ["ID", "Token", "Foydalanuvchi", "Telefon", "Tarif", "Holat", "Berilgan", "Aktivlashtirilgan", "Tugaydi",
                  "Qolgan kun", "Almashtirishlar", "Manba"],
                 [[r["id"], r["token"], r["full_name"], r["phone"], r["plan_title"], lab[licenses.eff_status(r)],
                   fmt_dt(r["created_at"]), fmt_dt(r["activated_at"]), fmt_dt(r["expires_at"]), licenses.days_left(r),
                   r["device_resets_used"], r["source"]] for r in rows])


def stats_xlsx(days: int = 90) -> bytes:
    rows = stats.daily(days)
    return _book("Kunlik statistika",
                 ["Sana", "Yangi foydalanuvchilar", "Buyurtmalar", "To'langan", "Tushum (so'm)", "Berilgan litsenziya",
                  "Aktivlashtirilgan"],
                 [[r["date"], r["users"], r["orders"], r["paid"], r["revenue"], r["issued"], r["activations"]] for r in rows])


KINDS = {"users": users_xlsx, "payments": payments_xlsx, "licenses": licenses_xlsx, "stats": stats_xlsx}
